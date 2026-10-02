"""Behavior tests for the rules that protect a project, a study, or the generated website.

Run with ``python manage.py test``. Database tests use an in-memory copy of the
project database, so /project/ is never touched.
"""

import json
import shutil
import socket
import subprocess
import time
import urllib.error
import urllib.request
from pathlib import Path
from tempfile import TemporaryDirectory

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from unittest import SkipTest, skipUnless

from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from PIL import Image

from .forms import FrameForm, ProjectSettingsForm, normalize_language, parse_extra_languages
from .models import Element, ElementLanguageOverride, Frame, ProjectSettings, Rule, Variable, name_key
from .services.components import component_map, default_color_css, default_font_css
from .services.content import add_language, append_content_row, create_content_workbook, load_content_table, rename_language, update_content_row
from .services.flow import STEP_X, create_frame, default_name, tidy_layout
from .services.frame_types import load_frame_types
from .services.generate import generate_dist, language_folder, reset_dist_directory
from .services.llm import parse_elements
from .services.log_keys import LockedLog, decrypt_event, encrypt_event, fingerprint, generate_key_pair, is_encrypted, load_private_key
from .services.project import image_size, list_materials, material_thumbnail, project_health, safe_child
from .services.stresstest import RUNTIME_FILES, build_check, device_checks, es5_problems
from .services.study_logs import SessionSummary, describe_visit, device_report, final_variables, frame_visits, import_jsonl, log_file_name, parse_events, readable_events, safe_name, session_kind
from .services.validate import validate_project
from .services.variables import check_name, check_value, kind_of, parse_value, rule_label, unknown_placeholders

ANSWER: str = '{"elements": [{"component": "basic-decision", "content_id": 1, "x": 960, "y": 540, "target": "fnr-2"}]}'
EVENT: dict[str, object] = {"participant_id": "R_abc", "visit_id": "20260910T101530Z-a1b2", "event": "frame", "seq": 1}


def make_project_folder(root: Path, languages: list[str]) -> ProjectSettings:
    """Files of a minimal project in ``root`` plus its settings row (test database)."""
    (root / "materials").mkdir(parents=True)
    (root / "logs").mkdir()
    (root / "project.sqlite3").touch()  # project_settings() checks that the file exists
    create_content_workbook(root / "content.xlsx", languages)
    (root / "default-colors.css").write_text(default_color_css(), encoding="utf-8")
    (root / "style-overrides.css").write_text("", encoding="utf-8")
    return ProjectSettings.objects.create(name="Demo", slug="demo", base_language=languages[0])


class LanguageTests(SimpleTestCase):
    def test_language_codes_are_normalized(self) -> None:
        self.assertEqual(normalize_language(" pt_br "), "pt-BR")

    def test_extra_languages_keep_order_without_duplicates(self) -> None:
        self.assertEqual(parse_extra_languages("nl-NL, DE-de, nl-nl", "en-US"), ["nl-NL", "de-DE"])

    def test_single_language_project_uses_the_dist_root(self) -> None:
        self.assertEqual(language_folder("en-US", ("en-US",)), "")

    def test_each_language_gets_its_own_dist_folder(self) -> None:
        self.assertEqual(language_folder("nl-NL", ("en-US", "nl-NL")), "nl-NL")


class FrameNameRuleTests(SimpleTestCase):
    def test_capitals_spaces_and_dashes_do_not_make_names_different(self) -> None:
        self.assertEqual(name_key("Frame 12"), name_key("frame-12"))

    def test_underscores_do_not_make_names_different(self) -> None:
        self.assertEqual(name_key("Frame_12"), name_key("frame 12"))

    def test_accents_do_not_make_names_different(self) -> None:
        self.assertEqual(name_key("Café scène"), name_key("cafe scene"))

    def test_default_name_repeats_the_number(self) -> None:
        self.assertEqual(default_name("Frame", 12, set()), "Frame 12")

    def test_default_name_skips_a_name_someone_already_chose(self) -> None:
        self.assertEqual(default_name("Frame", 12, {"frame-12"}), "Frame 13")


class LlmAnswerTests(SimpleTestCase):
    def parse(self, answer: str) -> list[dict[str, object]]:
        return parse_elements(answer, component_map(), content_ids={1}, targets={"fnr-2": 7})

    def test_answer_becomes_element_values(self) -> None:
        self.assertEqual(self.parse(ANSWER)[0]["target_frame_id"], 7)

    def test_missing_size_uses_the_component_default(self) -> None:
        self.assertEqual(self.parse(ANSWER)[0]["width"], component_map()["basic-decision"].default_width)

    def test_code_fences_around_the_answer_are_accepted(self) -> None:
        self.assertEqual(len(self.parse(f"```json\n{ANSWER}\n```")), 1)

    def test_unknown_component_names_the_element(self) -> None:
        with self.assertRaisesRegex(ValueError, "Element 1: unknown component"):
            self.parse('{"elements": [{"component": "robot", "x": 1, "y": 1}]}')

    def test_invented_content_id_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "content_id 99"):
            self.parse('{"elements": [{"component": "basic-speech-bubble", "content_id": 99, "x": 1, "y": 1}]}')

    def test_unknown_target_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "unknown target"):
            self.parse('{"elements": [{"component": "basic-decision", "x": 1, "y": 1, "target": "fnr-9"}]}')


class ContentWorkbookTests(SimpleTestCase):
    def test_created_workbook_puts_note_before_languages(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "content.xlsx"
            create_content_workbook(path, ["en-US", "nl-NL"])
            headers = [cell.value for cell in load_workbook(path).active[1]]
        self.assertEqual(headers, ["content_id", "note", "en-US", "nl-NL"])

    def test_appended_rows_get_the_next_id(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "content.xlsx"
            create_content_workbook(path, ["en-US"])
            append_content_row(path, "greeting", {"en-US": "Hello"})
            row = append_content_row(path, "reply", {"en-US": "Hi"})
            table = load_content_table(path)
        self.assertEqual(row.content_id, 2)
        self.assertEqual(table.by_id()[2].values["en-US"], "Hi")

    def test_a_row_can_be_changed_and_keeps_its_id(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "content.xlsx"
            create_content_workbook(path, ["en-US", "nl-NL"])
            append_content_row(path, "greeting", {"en-US": "Hello"})
            update_content_row(path, 1, "greeting (formal)", {"en-US": "Good day", "nl-NL": "Goedendag"})
            row = load_content_table(path).by_id()[1]
        self.assertEqual((row.note, row.values), ("greeting (formal)", {"en-US": "Good day", "nl-NL": "Goedendag"}))

    def test_changing_an_unknown_row_is_refused(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "content.xlsx"
            create_content_workbook(path, ["en-US"])
            with self.assertRaisesRegex(ValueError, "content_id 9"):
                update_content_row(path, 9, "", {"en-US": "x"})

    def test_line_breaks_typed_in_a_browser_are_kept_as_excel_line_breaks(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "content.xlsx"
            create_content_workbook(path, ["en-US"])
            append_content_row(path, "", {"en-US": "First line\r\nSecond line"})
            text = load_content_table(path).by_id()[1].values["en-US"]
        self.assertEqual(text, "First line\nSecond line")


class VariableRuleTests(SimpleTestCase):
    """The type of a variable comes from its initial value and every update or rule must fit it."""

    def test_numbers_and_text_are_told_apart(self) -> None:
        self.assertEqual([parse_value(raw) for raw in ("0", "2.5", "-3", "path 1", '"12"')], [0, 2.5, -3, "path 1", "12"])

    def test_the_kind_follows_the_value(self) -> None:
        self.assertEqual((kind_of(2.5), kind_of("go")), ("number", "text"))

    def test_adding_to_a_text_variable_is_refused(self) -> None:
        with self.assertRaisesRegex(ValueError, "Only numbers can be added to"):
            check_value("text", "1", "add")

    def test_setting_a_number_variable_to_text_is_refused(self) -> None:
        with self.assertRaisesRegex(ValueError, "must be a number"):
            check_value("number", "lots", "set")

    def test_setting_a_text_variable_to_a_number_needs_quotes(self) -> None:
        with self.assertRaisesRegex(ValueError, "in quotes"):
            check_value("text", "12", "set")

    def test_quoted_numbers_are_text(self) -> None:
        self.assertEqual(check_value("text", '"12"', "set"), "12")

    def test_less_than_only_compares_numbers(self) -> None:
        with self.assertRaisesRegex(ValueError, "only works with numbers"):
            check_value("text", "a", "<")

    def test_a_matching_rule_value_is_typed(self) -> None:
        self.assertEqual(check_value("number", "3", ">="), 3)

    def test_an_empty_value_is_refused(self) -> None:
        with self.assertRaises(ValueError):
            check_value("number", "  ", "set")

    def test_variable_names_are_identifiers(self) -> None:
        with self.assertRaises(ValueError):
            check_name("my score")

    def test_placeholders_that_are_not_variables_are_reported(self) -> None:
        self.assertEqual(unknown_placeholders("Score: {score} of {total}", {"score"}), ["total"])

    def test_rules_read_as_a_sentence(self) -> None:
        self.assertEqual((rule_label("score", ">=", "3"), rule_label(None, "==", "")), ("score is at least 3", "otherwise"))


class FontCssTests(SimpleTestCase):
    def test_every_component_gets_one_font_line(self) -> None:
        css = default_font_css()
        self.assertEqual(css.count("\n.component-"), len(component_map()))
        self.assertIn(".component-basic-narrator ", css)


class VisitHeaderTests(SimpleTestCase):
    """The first line of a log names the device and browser; the studio summarizes it."""

    def test_iphone_safari(self) -> None:
        header = {"user_agent": "Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.0 Mobile/15E148 Safari/604.1", "screen": {"width": 390, "height": 844}}
        self.assertEqual(describe_visit(header), "iPhone · Safari · 390×844")

    def test_android_model_and_chrome(self) -> None:
        header = {"user_agent": "Mozilla/5.0 (Linux; Android 13; SM-G991B) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Mobile Safari/537.36", "screen": {"width": 360, "height": 800}}
        self.assertEqual(describe_visit(header), "Android (SM-G991B) · Chrome · 360×800")

    def test_edge_on_windows_is_not_chrome(self) -> None:
        header = {"user_agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36 Edg/120.0.0.0"}
        self.assertEqual(describe_visit(header), "Windows · Edge")

    def test_a_visit_line_in_the_timeline(self) -> None:
        events = [{"event": "visit", "local_time": "2026-09-15T10:41:02+02:00", "user_agent": "iPad", "timestamp": "2026-09-15T08:41:02Z"}]
        self.assertEqual(readable_events(events, {}, {}), ["Visit started at 2026-09-15T10:41:02+02:00 on iPad"])

    def test_the_visit_line_names_the_tiltale_version_when_logged(self) -> None:
        events = [{"event": "visit", "local_time": "2026-09-17T10:41:02+02:00", "user_agent": "iPad", "tiltale_version": "2.3.0"}]
        self.assertEqual(readable_events(events, {}, {}), ["Visit started at 2026-09-17T10:41:02+02:00 on iPad (story made with TilTale 2.3.0)"])


class LogKeyTests(SimpleTestCase):
    """Encrypted log lines (as log.php writes them) can only be read with the project's private key."""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        cls.private_pem, cls.public_pem = generate_key_pair()

    def test_a_line_round_trips(self) -> None:
        event = {"participant_id": "R_1", "visit_id": "v1", "event": "frame", "frame": "fnr-1", "text": "héllo"}
        line = encrypt_event(event, self.public_pem)
        self.assertEqual(sorted(json.loads(line)), ["data", "enc", "iv", "tag"])
        private = load_private_key(self.private_pem.encode(), self.public_pem)
        self.assertEqual(parse_events(line, "x", private), [event])

    def test_reading_without_the_key_is_refused(self) -> None:
        line = encrypt_event({"event": "frame"}, self.public_pem)
        with self.assertRaises(LockedLog):
            parse_events(line, "x")

    def test_the_wrong_key_is_refused(self) -> None:
        other_private, _other_public = generate_key_pair()
        with self.assertRaisesRegex(ValueError, "another project"):
            load_private_key(other_private.encode(), self.public_pem)
        with self.assertRaisesRegex(ValueError, "not a key file"):
            load_private_key(b"hello", self.public_pem)

    def test_a_tampered_line_is_refused(self) -> None:
        record = json.loads(encrypt_event({"event": "frame"}, self.public_pem))
        record["data"] = record["data"][:-4] + "AAAA"
        private = load_private_key(self.private_pem.encode(), self.public_pem)
        with self.assertRaises(ValueError):
            decrypt_event(record, private)

    def test_plain_and_encrypted_lines_can_share_a_file(self) -> None:
        private = load_private_key(self.private_pem.encode(), self.public_pem)
        text = json.dumps({"event": "a"}) + "\n" + encrypt_event({"event": "b"}, self.public_pem)
        self.assertEqual([e["event"] for e in parse_events(text, "x", private)], ["a", "b"])

    def test_fingerprints_are_short_and_stable(self) -> None:
        self.assertEqual(fingerprint(self.public_pem), fingerprint(self.public_pem))
        self.assertRegex(fingerprint(self.public_pem), r"^([0-9a-f]{4}:){7}[0-9a-f]{4}$")


NASTY_EVENTS: list[dict] = [  # everything a story could ever send
    {"participant_id": "P1", "visit_id": "v1", "event": "frame", "frame": "start", "seq": 1},
    {"participant_id": "P1", "visit_id": "v1", "event": "choice", "seq": 2,
     "text": "Café naïve — “quotes” ‘and’ \\ / <b>&amp;</b> 😀 中文 🇳🇱 \u0000 \t"},
    {"participant_id": "P1", "visit_id": "v1", "event": "variables", "seq": 3,
     "data": {"nested": [1, 2.5, None, True, {"k": "v"}], "empty": "", "big": 10 ** 18}},
    {"participant_id": "P1", "visit_id": "v1", "event": "note", "seq": 4, "long": "x" * 30000},
]


@skipUnless(shutil.which("php"), "PHP is not installed")
class LogPhpEncryptionTests(SimpleTestCase):
    """What log.php writes with log-key.pem next to it must open with the key file, and nothing else."""

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        cls.private_pem, cls.public_pem = generate_key_pair()
        cls.temp = TemporaryDirectory()
        cls.dist = Path(cls.temp.name)
        shutil.copy2(Path(settings.BASE_DIR) / "runtime" / "log.php", cls.dist / "log.php")
        (cls.dist / "log-key.pem").write_text(cls.public_pem, encoding="utf-8")
        if subprocess.run(["php", "-r", "exit(extension_loaded('openssl') ? 0 : 1);"]).returncode:
            raise SkipTest("PHP has no OpenSSL extension")
        with socket.socket() as probe:  # a free port
            probe.bind(("127.0.0.1", 0))
            cls.port = probe.getsockname()[1]
        cls.server = subprocess.Popen(["php", "-S", f"127.0.0.1:{cls.port}", "-t", str(cls.dist)],
                                      stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for _ in range(50):  # wait until PHP listens
            try:
                socket.create_connection(("127.0.0.1", cls.port), timeout=0.2).close()
                break
            except OSError:
                time.sleep(0.1)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.server.terminate()
        cls.server.wait()
        cls.temp.cleanup()
        super().tearDownClass()

    def setUp(self) -> None:
        shutil.rmtree(self.dist / "logs", ignore_errors=True)  # every test starts with an empty logs/ folder

    def post(self, payload: object) -> int:
        request = urllib.request.Request(f"http://127.0.0.1:{self.port}/log.php", data=json.dumps(payload).encode(),
                                         headers={"Content-Type": "application/json"}, method="POST")
        try:
            return urllib.request.urlopen(request, timeout=10).status
        except urllib.error.HTTPError as error:
            return error.code

    def log_lines(self) -> list[str]:
        return (self.dist / "logs" / "P1--v1.jsonl").read_text(encoding="utf-8").splitlines()

    def test_php_encrypts_and_python_decrypts_every_kind_of_event(self) -> None:
        self.assertEqual([self.post(event) for event in NASTY_EVENTS[:2]], [204, 204])  # one per request
        self.assertEqual(self.post(NASTY_EVENTS[2:]), 204)  # a batch, as tiltale.js sends it
        lines = self.log_lines()
        self.assertEqual(len(lines), len(NASTY_EVENTS))
        private = load_private_key(self.private_pem.encode(), self.public_pem)
        for line, original in zip(lines, NASTY_EVENTS):
            self.assertTrue(is_encrypted(json.loads(line)))
            for secret in ("participant_id", "Café", "xxxx", "start"):
                self.assertNotIn(secret, line, "plaintext leaked into the log")
            event = decrypt_event(json.loads(line), private)
            self.assertIn("received_at", event)
            event.pop("received_at")
            self.assertEqual(event, original)
        # and through the studio's own reader, the way Results opens a file
        self.assertEqual(len(parse_events("\n".join(lines), "P1--v1.jsonl", private)), len(NASTY_EVENTS))

    def test_the_wrong_key_cannot_read_php_output(self) -> None:
        self.post(NASTY_EVENTS[0])
        other_private, other_public = generate_key_pair()
        other = load_private_key(other_private.encode(), other_public)
        with self.assertRaises(ValueError):
            decrypt_event(json.loads(self.log_lines()[0]), other)

    def test_a_key_file_survives_being_saved_on_any_os(self) -> None:
        """Windows line endings, a BOM, a trailing blank line: the file still opens the logs."""
        self.post(NASTY_EVENTS[0])
        record = json.loads(self.log_lines()[0])
        for variant in (self.private_pem.replace("\n", "\r\n"), "﻿" + self.private_pem, self.private_pem + "\n\n"):
            private = load_private_key(variant.encode("utf-8"), self.public_pem)
            self.assertEqual(decrypt_event(record, private)["event"], "frame")

    def test_without_the_key_file_php_writes_plain_lines(self) -> None:
        (self.dist / "log-key.pem").rename(self.dist / "log-key.pem.off")
        try:
            self.assertEqual(self.post(NASTY_EVENTS[0]), 204)
            self.assertFalse(is_encrypted(json.loads(self.log_lines()[-1])))
        finally:
            (self.dist / "log-key.pem.off").rename(self.dist / "log-key.pem")

    def test_a_damaged_key_file_writes_nothing_rather_than_plaintext(self) -> None:
        (self.dist / "log-key.pem").write_text("-----BEGIN PUBLIC KEY-----\ngarbage\n-----END PUBLIC KEY-----\n")
        try:
            self.assertEqual(self.post(NASTY_EVENTS[0]), 500)
            self.assertFalse((self.dist / "logs" / "P1--v1.jsonl").exists())
        finally:
            (self.dist / "log-key.pem").write_text(self.public_pem, encoding="utf-8")


class FrameTypeTests(SimpleTestCase):
    def test_every_kind_has_a_catalog_entry_with_words(self) -> None:
        types = load_frame_types()
        self.assertEqual(list(types), list(Frame.Kind.values))
        self.assertTrue(all(item.name and item.description and item.help for item in types.values()))


class SettingsFormTests(SimpleTestCase):
    def form(self, **changes: object) -> ProjectSettingsForm:
        data = {
            "name": "Demo", "frame_width": "1920", "frame_height": "1080", "default_delay_seconds": "0.75",
            "letterbox_color": "#000000", "participant_parameter": "ppn", "finish_redirect_url": "",
        }
        return ProjectSettingsForm({**data, **changes})

    def test_redirect_url_may_contain_the_id_placeholder(self) -> None:
        self.assertTrue(self.form(finish_redirect_url="https://x.qualtrics.com/jfe/form/SV_1?ppn={ID}").is_valid())

    def test_participant_parameter_rejects_spaces(self) -> None:
        self.assertFalse(self.form(participant_parameter="my id").is_valid())

    def test_frame_size_accepts_decimals(self) -> None:
        self.assertTrue(self.form(frame_width="1080.5").is_valid())


class FileBoundaryTests(SimpleTestCase):
    def test_safe_child_rejects_parent_traversal(self) -> None:
        with TemporaryDirectory() as directory, self.assertRaises(ValueError):
            safe_child(Path(directory), "../outside.txt")

    def test_regeneration_refuses_a_dist_that_contains_the_project(self) -> None:
        with TemporaryDirectory() as directory:
            base = Path(directory)
            with override_settings(BASE_DIR=base, PROJECT_DIR=base / "data" / "project", DIST_DIR=base / "data"):
                with self.assertRaises(ValueError):
                    reset_dist_directory()

    def test_regeneration_keeps_downloaded_study_logs(self) -> None:
        with TemporaryDirectory() as directory:
            base = Path(directory)
            (base / "dist" / "logs").mkdir(parents=True)
            (base / "dist" / "logs" / "R_abc--v1.jsonl").write_text("{}\n", encoding="utf-8")
            (base / "dist" / "old.html").write_text("", encoding="utf-8")
            with override_settings(BASE_DIR=base, PROJECT_DIR=base / "project", DIST_DIR=base / "dist"):
                reset_dist_directory()
            names = sorted(path.name for path in (base / "dist").rglob("*"))
        self.assertEqual(names, ["R_abc--v1.jsonl", "logs"])


class ImageSizeTests(SimpleTestCase):
    def test_exif_rotated_photos_report_their_upright_size(self) -> None:
        with TemporaryDirectory() as directory:
            (Path(directory) / "materials").mkdir()
            exif = Image.Exif()
            exif[0x0112] = 6  # Orientation: rotate 90° when shown
            Image.new("RGB", (400, 300)).save(Path(directory) / "materials" / "photo.jpg", exif=exif)
            with override_settings(PROJECT_DIR=Path(directory)):
                self.assertEqual(image_size("photo.jpg"), (300, 400))


class StudyLogTests(SimpleTestCase):
    def test_every_visit_gets_its_own_log_file(self) -> None:
        self.assertEqual(log_file_name(EVENT), "R_abc--20260910T101530Z-a1b2.jsonl")

    def test_unsafe_participant_ids_are_made_file_safe(self) -> None:
        self.assertEqual(safe_name("../R abc"), "R-abc")

    def test_play_test_visits_are_recognized(self) -> None:
        self.assertEqual(session_kind("playtest-en-us-123"), "play-test")

    def test_qualtrics_ids_count_as_study_visits(self) -> None:
        self.assertEqual(session_kind("R_1abcDEF"), "study")

    def test_seconds_per_frame_run_until_the_next_frame_or_the_end(self) -> None:
        events = [
            {"event": "frame", "frame": "fnr-1", "timestamp": "2026-09-10T10:00:00Z"},
            {"event": "choice", "frame": "fnr-1", "element_id": 5, "timestamp": "2026-09-10T10:00:12Z"},
            {"event": "frame", "frame": "fnr-2", "timestamp": "2026-09-10T10:00:12Z"},
            {"event": "Story finished", "frame": "fnr-2", "timestamp": "2026-09-10T10:00:15Z"},
        ]
        self.assertEqual(frame_visits(events), [("fnr-1", 12.0), ("fnr-2", 3.0)])

    def test_readable_lines_use_frame_and_element_names(self) -> None:
        events = [
            {"event": "frame", "frame": "fnr-1", "timestamp": "2026-09-10T10:00:00Z", "how": "resumed"},
            {"event": "choice", "frame": "fnr-1", "element_id": 5, "timestamp": "2026-09-10T10:00:12Z"},
        ]
        lines = readable_events(events, {"fnr-1": "Intro"}, {5: "Choice button (Go on)"})
        self.assertEqual(lines, ["IDN refreshed", "Intro: visited (last frame)", "Intro: clicked 'Choice button (Go on)'"])

    def test_variable_changes_and_decisions_get_their_own_lines(self) -> None:
        events = [
            {"event": "frame", "frame": "fnr-1", "timestamp": "2026-09-10T10:00:00Z"},
            {"event": "variable", "frame": "fnr-1", "variable": "score", "from": 0, "to": 1, "timestamp": "2026-09-10T10:00:05Z"},
            {"event": "decision", "frame": "fnr-2", "element_id": "rule-7", "target": "fnr-3", "timestamp": "2026-09-10T10:00:05Z"},
            {"event": "frame", "frame": "fnr-3", "timestamp": "2026-09-10T10:00:05Z"},
        ]
        names = {"fnr-1": "Intro", "fnr-2": "Check", "fnr-3": "Good end"}
        lines = readable_events(events, names, {"rule-7": "score is at least 1"})
        self.assertEqual(lines, ["Intro: visited for 5 s", "score: 0 → 1", "Check: score is at least 1 → Good end", "Good end: visited (last frame)"])
        self.assertEqual(frame_visits(events), [("fnr-1", 5.0), ("fnr-2", None), ("fnr-3", None)])
        self.assertEqual(final_variables(events), {"score": 1})

    def test_a_javascript_error_gets_a_readable_line(self) -> None:
        events = [{"event": "error", "message": "x is not defined", "file": "https://example.org/tiltale.js", "line": 212}]
        self.assertEqual(readable_events(events, {}, {}), ["JavaScript error: x is not defined (https://example.org/tiltale.js, line 212)"])

    def test_import_never_overwrites_an_existing_log(self) -> None:
        data = (json.dumps(EVENT) + "\n").encode()
        with TemporaryDirectory() as directory:
            (Path(directory) / "logs").mkdir()
            with override_settings(PROJECT_DIR=Path(directory)):
                import_jsonl("first.jsonl", data)
                with self.assertRaises(ValueError):
                    import_jsonl("again.jsonl", data)

    def test_import_rejects_two_visits_in_one_file(self) -> None:
        data = (json.dumps(EVENT) + "\n" + json.dumps({**EVENT, "visit_id": "other"}) + "\n").encode()
        with TemporaryDirectory() as directory:
            (Path(directory) / "logs").mkdir()
            with override_settings(PROJECT_DIR=Path(directory)), self.assertRaises(ValueError):
                import_jsonl("mixed.jsonl", data)


def make_summary(device: str, events: list[dict[str, object]], finished: bool = False, kind: str = "study") -> SessionSummary:
    return SessionSummary(file_name="x.jsonl", participant_id="p", visit_id="v", language="en", event_count=len(events), started="",
                          last_frame="", finished=finished, kind=kind, device=device, local_time="", events=tuple(events))


class DeviceReportTests(SimpleTestCase):
    """The Results page sums up every device the story was opened on: visits, finished visits and errors."""

    def test_visits_are_counted_per_device(self) -> None:
        rows = device_report([
            make_summary("iPhone · Safari", [{"event": "visit"}], finished=True),
            make_summary("iPhone · Safari", [{"event": "visit"}]),
            make_summary("Android · Chrome", [{"event": "visit"}, {"event": "error", "message": "boom"}]),
        ])
        self.assertEqual([(row["device"], row["visits"], row["finished"], row["errors"]) for row in rows],
                         [("Android · Chrome", 1, 0, 1), ("iPhone · Safari", 2, 1, 0)])  # problems first

    def test_a_browser_too_old_is_its_own_column(self) -> None:
        rows = device_report([make_summary("Android · Chrome", [{"event": "Browser not supported", "missing": ["fetch"]}])])
        self.assertEqual((rows[0]["unsupported"], rows[0]["visits"]), (1, 1))

    def test_a_visit_without_a_header_is_an_unknown_device(self) -> None:
        self.assertEqual(device_report([make_summary("", [{"event": "frame"}])])[0]["device"], "Unknown device")


class StressTestChecklistTests(SimpleTestCase):
    """The Stress test page's rows that need no browser (services/stresstest.py)."""

    def test_the_story_player_stays_es5(self) -> None:
        """Old phones run only ES5; python manage.py test is where a developer hears about a slip first."""
        for name in RUNTIME_FILES:
            self.assertEqual(es5_problems((settings.RUNTIME_DIR / name).read_text(encoding="utf-8"), name), [])

    def test_modern_syntax_is_named_with_its_line(self) -> None:
        text = 'var a = 1; // a const in a comment does not count\nvar s = "let it be";\nconst b = () => `x`;'
        self.assertEqual(es5_problems(text, "x.js"), ["x.js line 3: 'const'", "x.js line 3: '=>'", "x.js line 3: '`'", "x.js line 3: '`'"])

    def test_regenerate_warnings_make_an_orange_row(self) -> None:
        report = {"issues": [{"severity": "warning", "message": "Intro: text is about 8.0px on a iPhone 5s / SE 1st gen (minimum 12px)."}]}
        self.assertEqual((build_check(report).status, build_check({"issues": []}).status, build_check(None).status), ("orange", "green", "red"))

    def test_real_phones_are_counted_from_study_visits_only(self) -> None:
        visits = [
            make_summary("iPhone · Safari", [{"event": "visit"}], finished=True),
            make_summary("Android · Chrome", [{"event": "visit"}, {"event": "error", "message": "boom"}]),
        ]
        studio_run = make_summary("Windows · Chrome", [{"event": "visit"}], finished=True, kind="play-test")
        phones, errors = device_checks([*visits, studio_run])
        self.assertEqual((phones.status, phones.message), ("orange", "1 so far: iPhone · Safari"))
        self.assertEqual((errors.status, errors.message), ("red", "Errors or a browser too old on: Android · Chrome"))
        self.assertEqual(device_checks([])[0].status, "red")


class PreviewTests(SimpleTestCase):
    def test_missing_preview_page_explains_what_to_do(self) -> None:
        response = self.client.get(reverse("studio:preview_file", kwargs={"path": "not-built/index.html"}))
        self.assertContains(response, "Regenerate", status_code=404)

    def test_branding_serves_no_other_repository_files(self) -> None:
        response = self.client.get(reverse("studio:branding", kwargs={"name": "manage.py"}))
        self.assertEqual(response.status_code, 404)


class FrameNameTests(TestCase):
    databases = {"project"}

    def form(self, frame: Frame, name: str) -> FrameForm:
        data = {"name": name, "background_type": "none", "background_color": "#111111", "background_image": ""}
        return FrameForm(data, instance=frame, materials=[], content_ids=set())

    def test_frames_can_be_renamed_with_spaces(self) -> None:
        frame = Frame.objects.create(name="frame-1")
        self.assertTrue(self.form(frame, "Intro scene").is_valid())

    def test_names_with_the_same_code_name_are_rejected(self) -> None:
        Frame.objects.create(name="Frame-12")
        frame = Frame.objects.create(name="frame-13")
        self.assertIn("name", self.form(frame, "frame 12").errors)

    def test_renaming_to_a_variant_of_its_own_name_is_allowed(self) -> None:
        frame = Frame.objects.create(name="frame-1")
        self.assertTrue(self.form(frame, "Frame 1").is_valid())

    def test_name_without_letters_or_digits_is_rejected(self) -> None:
        frame = Frame.objects.create(name="frame-1")
        self.assertIn("name", self.form(frame, "!!").errors)

    def test_new_frame_name_repeats_its_id_number(self) -> None:
        frame = create_frame()
        self.assertEqual((frame.key, frame.name), (f"fnr-{frame.pk}", f"Frame {frame.pk}"))

    def test_new_picker_frames_are_called_picker(self) -> None:
        frame = create_frame("picker")
        self.assertEqual(frame.name, f"Picker {frame.pk}")

    def test_fade_is_one_three_way_choice(self) -> None:
        """One of: no fade, crossfade, from black. "From black" alone must still fade (it once did not)."""
        frame = Frame.objects.create(name="frame-1")
        for choice, fade_in, fade_from_black in (("black", True, True), ("previous", True, False), ("none", False, False)):
            data = {"name": "frame-1", "fade": choice, "background_type": "none", "background_color": "#111111", "background_image": ""}
            FrameForm(data, instance=frame, materials=[], content_ids=set()).save()
            frame.refresh_from_db()
            self.assertEqual((frame.fade_in, frame.fade_from_black), (fade_in, fade_from_black), choice)

    def test_a_new_background_image_fills_the_frame_again(self) -> None:
        frame = Frame.objects.create(name="frame-1", background_type="image", background_image="old.png", background_width=50)
        data = {"name": "frame-1", "background_type": "image", "background_color": "#111111", "background_image": "new.png"}
        FrameForm(data, instance=frame, materials=["new.png"], content_ids=set()).save()
        self.assertIsNone(Frame.objects.get(pk=frame.pk).background_box)

    def test_renaming_keeps_the_id(self) -> None:
        frame = create_frame()
        form = self.form(frame, "Intro scene")
        form.save()
        self.assertEqual(Frame.objects.get(name="Intro scene").key, f"fnr-{frame.pk}")


class FlowchartTests(TestCase):
    databases = {"project"}

    def test_new_linked_frame_is_placed_right_of_its_source(self) -> None:
        source = Frame.objects.create(name="frame-1", flow_x=100, flow_y=50)
        frame = create_frame(linked_from=source)
        self.assertEqual((frame.flow_x, frame.flow_y), (100 + STEP_X, 50))

    def test_second_frame_from_the_same_source_does_not_overlap(self) -> None:
        source = Frame.objects.create(name="frame-1", flow_x=100, flow_y=50)
        first = create_frame(linked_from=source)
        second = create_frame(linked_from=source)
        self.assertNotEqual((first.flow_x, first.flow_y), (second.flow_x, second.flow_y))


class ProjectTestCase(TestCase):
    """A throw-away project folder, dist folder and in-memory project database."""

    databases = {"project"}
    languages: list[str] = ["en-US"]

    def setUp(self) -> None:
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        project_dir = self.root / "project"
        settings_override = override_settings(
            BASE_DIR=self.root, PROJECT_DIR=project_dir, PROJECT_DB=project_dir / "project.sqlite3", DIST_DIR=self.root / "dist",
        )
        settings_override.enable()
        self.addCleanup(settings_override.disable)
        self.project = make_project_folder(project_dir, self.languages)


class StudioViewTests(ProjectTestCase):

    def test_the_preview_logger_takes_one_event_or_a_batch(self) -> None:
        """tiltale.js sends up to 25 queued events in one request; log.php and this stand-in accept both shapes."""
        url = reverse("studio:preview_log")
        one = {**EVENT, "seq": 1}
        batch = [{**EVENT, "seq": 2}, {**EVENT, "seq": 3, "event": "Story finished"}]
        self.assertEqual(self.client.post(url, json.dumps(one), content_type="application/json").status_code, 204)
        self.assertEqual(self.client.post(url, json.dumps(batch), content_type="application/json").status_code, 204)
        self.assertEqual(self.client.post(url, json.dumps([1, 2]), content_type="application/json").status_code, 400)
        lines = (self.root / "project" / "logs" / "R_abc--20260910T101530Z-a1b2.jsonl").read_text(encoding="utf-8").splitlines()
        self.assertEqual([json.loads(line)["seq"] for line in lines], [1, 2, 3])

    def test_the_test_page_shows_the_stress_checklist(self) -> None:
        (self.root / "dist").mkdir()
        (self.root / "dist" / ".build.json").write_text(json.dumps({"builds": [{"label": "en", "folder": ""}], "source_timestamp": 0, "issues": []}), encoding="utf-8")
        response = self.client.get(reverse("studio:test"))
        self.assertContains(response, 'data-stress="lost-connection"')
        self.assertContains(response, 'data-result="red" data-result-message="No finished visit from a real phone yet')

    def test_old_playtest_and_stresstest_links_land_on_the_test_page(self) -> None:
        for name in ("playtest", "stresstest"):
            self.assertEqual(self.client.get(reverse(f"studio:{name}")).url, reverse("studio:test"))

    def test_dragged_flowchart_positions_are_stored(self) -> None:
        first = Frame.objects.create(name="frame-1")
        second = Frame.objects.create(name="frame-2")
        body = {"positions": [{"id": first.id, "x": 10, "y": 20}, {"id": second.id, "x": 300, "y": 20}]}
        self.client.post(reverse("studio:flow_positions_api"), json.dumps(body), content_type="application/json")
        positions = list(Frame.objects.values_list("name", "flow_x", "flow_y"))
        self.assertEqual(positions, [("frame-1", 10.0, 20.0), ("frame-2", 300.0, 20.0)])

    def test_deleting_logs_or_materials_does_not_block_the_studio(self) -> None:
        """They hold nothing of their own, so TilTale makes them again instead of asking for a backup."""
        for name in ("logs", "materials"):
            shutil.rmtree(self.root / "project" / name)
        self.assertEqual(project_health().problems, ())
        self.assertTrue((self.root / "project" / "logs").is_dir())

    def test_new_elements_go_in_front_and_the_list_can_be_restacked(self) -> None:
        frame = Frame.objects.create(name="frame-1")
        first = Element.objects.create(frame=frame, component="basic-decision")
        second = Element.objects.create(frame=frame, component="basic-decision")
        self.assertEqual(list(frame.elements.values_list("id", flat=True)), [first.id, second.id])
        url = reverse("studio:element_order_api", args=[frame.id])
        body = {"order": [first.id, second.id]}  # front to back: first now in front
        self.client.post(url, json.dumps(body), content_type="application/json")
        self.assertEqual(list(frame.elements.values_list("id", flat=True)), [second.id, first.id])
        stale = self.client.post(url, json.dumps({"order": [first.id]}), content_type="application/json")
        self.assertEqual(stale.status_code, 409)  # the list no longer matches the frame

    def test_target_list_is_alphabetical_and_skips_the_frame_itself(self) -> None:
        frame = Frame.objects.create(name="Middle")
        Frame.objects.create(name="zebra")
        Frame.objects.create(name="Apple")
        Element.objects.create(frame=frame, component="basic-decision")
        html = self.client.get(reverse("studio:frame_editor", args=[frame.id])).content.decode()
        self.assertLess(html.index("Go to Apple"), html.index("Go to zebra"))
        self.assertNotIn("Go to Middle", html)

    def test_single_language_projects_cannot_create_picker_frames(self) -> None:
        self.client.post(reverse("studio:new_frame"), {"kind": "picker"})
        self.assertFalse(Frame.objects.filter(kind=Frame.Kind.PICKER).exists())

    def test_new_frames_get_their_kind_and_a_matching_name(self) -> None:
        self.client.post(reverse("studio:new_frame"), {"kind": "validation"})
        frame = Frame.objects.get()
        self.assertEqual((frame.kind, frame.name), ("validation", f"Validation {frame.pk}"))

    def test_unknown_kinds_are_refused(self) -> None:
        self.client.post(reverse("studio:new_frame"), {"kind": "poster"})
        self.assertFalse(Frame.objects.exists())

    def test_regenerate_works_twice_in_a_row(self) -> None:
        Frame.objects.create(name="frame-1")
        generate_dist(self.project)
        generate_dist(self.project)
        self.assertTrue((self.root / "dist" / "index.html").is_file())

    def test_single_language_story_is_the_dist_root(self) -> None:
        Frame.objects.create(name="frame-1")
        report = generate_dist(self.project)
        self.assertEqual([build.folder for build in report.builds], [""])

    def test_adding_a_frame_does_not_rebuild_dist(self) -> None:
        """Regenerate can take minutes on a big story, so only the Regenerate button starts it."""
        self.client.post(reverse("studio:new_frame"))
        self.assertFalse((self.root / "dist").exists())

    def test_story_uses_frame_ids_not_names(self) -> None:
        frame = Frame.objects.create(name="Intro scene")
        generate_dist(self.project)
        self.assertIn(f'"start_frame":"{frame.key}"', (self.root / "dist" / "story.js").read_text(encoding="utf-8"))

    def test_generated_pages_carry_the_version_and_a_restart_route(self) -> None:
        Frame.objects.create(name="frame-1", fade_in=True, fade_from_black=True)
        generate_dist(self.project)
        self.assertIn("Made with TilTale version", (self.root / "dist" / "index.html").read_text(encoding="utf-8"))
        story = (self.root / "dist" / "story.js").read_text(encoding="utf-8")
        self.assertIn(f'"version":"{settings.TILTALE_VERSION}"', story)  # tiltale.js logs it in every visit header
        self.assertIn('"fade_in":true,"fade_from_black":true', story)
        self.assertTrue((self.root / "dist" / "restart" / "index.html").is_file())

    def test_documents_get_their_image_and_close_text_in_the_story(self) -> None:
        append_content_row(self.root / "project" / "content.xlsx", "", {"en-US": "Close"})
        Image.new("RGB", (400, 200)).save(self.root / "project" / "materials" / "poster.png")
        Frame.objects.create(name="frame-1")
        Frame.objects.create(name="poster", kind="document", background_type="image", background_image="poster.png", close_content_id=1)
        generate_dist(self.project)
        story = (self.root / "dist" / "story.js").read_text(encoding="utf-8")
        self.assertIn('"kind":"document"', story)
        self.assertIn('"height":200,"path":"assets/poster-', story)
        self.assertIn('"close_text":"Close"', story)

    def test_variables_rules_updates_and_restarts_are_in_the_story(self) -> None:
        score = Variable.objects.create(name="score", initial_value="0")
        first = Frame.objects.create(name="frame-1")
        check = Frame.objects.create(name="check", kind="validation")
        good = Frame.objects.create(name="good")
        Element.objects.create(frame=first, component="basic-decision", target_frame=check, update_variable=score, update_operation="add", update_value="1")
        Element.objects.create(frame=good, component="basic-decision", restarts_story=True)
        Rule.objects.create(frame=check, order=1, variable=score, comparator=">=", value="1", target_frame=good)
        Rule.objects.create(frame=check, order=2, target_frame=first)
        generate_dist(self.project)
        story = (self.root / "dist" / "story.js").read_text(encoding="utf-8")
        self.assertIn('"variables":{"score":0}', story)
        self.assertIn('"update":{"variable":"score","operation":"add","value":1}', story)
        self.assertIn(f'"label":"score is at least 1","variable":"score","comparator":">=","value":1,"target":"{good.key}"', story)
        self.assertIn('"label":"otherwise","variable":null', story)
        self.assertIn('"restarts_story":true', story)
        self.assertIn('"start_page":"index.html"', story)

    def test_a_scoreboard_on_every_frame_skips_the_frames_where_it_is_hidden(self) -> None:
        shown = Frame.objects.create(name="frame-1")
        hidden = Frame.objects.create(name="frame-2")
        board = Element.objects.create(frame=None, component="scoreboard-pill")
        board.hidden_on.add(hidden)
        generate_dist(self.project)
        story = json.loads((self.root / "dist" / "story.js").read_text(encoding="utf-8").split("var STORY = ")[1].rstrip(";\n"))
        elements = {frame["name"]: [element["id"] for element in frame["elements"]] for frame in story["frames"]}
        self.assertEqual((elements[shown.key], elements[hidden.key]), ([board.id], []))

    def test_saving_a_type_mismatch_on_an_element_is_refused(self) -> None:
        score = Variable.objects.create(name="score", initial_value="0")
        element = Element.objects.create(frame=Frame.objects.create(name="frame-1"), component="basic-decision")
        data = {
            "x": "1", "y": "1", "width": "100", "height": "50", "font_size": "30", "delay_mode": "none",
            "update_variable": str(score.id), "update_operation": "set", "update_value": "lots",
        }
        self.client.post(reverse("studio:save_element", args=[element.id]), data)
        self.assertIsNone(Element.objects.get(pk=element.id).update_variable)

    def test_rules_are_saved_in_order_with_the_otherwise_row_last(self) -> None:
        score = Variable.objects.create(name="score", initial_value="0")
        check = Frame.objects.create(name="check", kind="validation")
        good = Frame.objects.create(name="good")
        data = {
            "rule.0.id": "", "rule.0.variable": str(score.id), "rule.0.comparator": ">=", "rule.0.value": "3", "rule.0.target": f"frame:{good.id}",
            "else.id": "", "else.target": "new",
        }
        self.client.post(reverse("studio:save_rules", args=[check.id]), data)
        rules = list(check.rules.all())
        self.assertEqual([(rule.variable_id, rule.comparator, rule.value) for rule in rules], [(score.id, ">=", "3"), (None, "==", "")])
        self.assertEqual(rules[1].target_frame.name, f"Frame {rules[1].target_frame.pk}")

    def test_a_rule_with_the_wrong_type_is_refused(self) -> None:
        path = Variable.objects.create(name="path", initial_value="a")
        check = Frame.objects.create(name="check", kind="validation")
        data = {"rule.0.id": "", "rule.0.variable": str(path.id), "rule.0.comparator": "<", "rule.0.value": "b", "rule.0.target": "", "else.id": "", "else.target": ""}
        self.client.post(reverse("studio:save_rules", args=[check.id]), data)
        self.assertFalse(check.rules.exists())

    def test_a_variable_cannot_change_type_while_a_rule_uses_the_old_one(self) -> None:
        score = Variable.objects.create(name="score", initial_value="0")
        Rule.objects.create(frame=Frame.objects.create(name="check", kind="validation"), variable=score, comparator=">=", value="3")
        self.client.post(reverse("studio:save_variables"), {f"name.{score.id}": "score", f"initial.{score.id}": "none"})
        self.assertEqual(Variable.objects.get(pk=score.id).initial_value, "0")

    def test_a_used_variable_cannot_be_deleted(self) -> None:
        score = Variable.objects.create(name="score", initial_value="0")
        Element.objects.create(frame=Frame.objects.create(name="frame-1"), component="basic-decision", update_variable=score, update_value="1")
        self.client.post(reverse("studio:delete_variable", args=[score.id]))
        self.assertTrue(Variable.objects.filter(pk=score.id).exists())

    def test_validation_points_without_an_otherwise_row_are_reported(self) -> None:
        Frame.objects.create(name="frame-1")
        score = Variable.objects.create(name="score", initial_value="0")
        Rule.objects.create(frame=Frame.objects.create(name="check", kind="validation"), variable=score, comparator=">=", value="3")
        issues = validate_project(self.project, load_content_table(self.root / "project" / "content.xlsx"))
        self.assertTrue(any("otherwise" in issue.message for issue in issues))

    def test_generating_a_key_closes_the_studio_until_the_key_is_saved(self) -> None:
        self.client.post(reverse("studio:generate_log_key"))
        project = ProjectSettings.objects.get()
        self.assertTrue(project.logs_protected and project.log_key_pending.startswith("-----BEGIN PRIVATE KEY-----"))
        self.assertRedirects(self.client.get(reverse("studio:develop")), reverse("studio:log_key"), fetch_redirect_response=False)
        download = self.client.get(reverse("studio:download_log_key"))
        self.assertIn(f'filename="{project.slug}-log-key.pem"', download["Content-Disposition"])
        self.client.post(reverse("studio:confirm_log_key"), {})  # without the file: still pending
        self.assertTrue(ProjectSettings.objects.get().log_key_pending)
        self.client.post(reverse("studio:confirm_log_key"), {"key_file": SimpleUploadedFile("key.pem", b"not a key")})
        self.assertTrue(ProjectSettings.objects.get().log_key_pending)  # a broken download: still pending
        self.client.post(reverse("studio:confirm_log_key"), {"key_file": SimpleUploadedFile("key.pem", download.content)})
        self.assertEqual(ProjectSettings.objects.get().log_key_pending, "")
        self.assertEqual(self.client.get(reverse("studio:develop")).status_code, 200)
        self.assertEqual(self.client.post(reverse("studio:generate_log_key")).status_code, 302)  # once: the key stays
        self.assertEqual(ProjectSettings.objects.get().log_public_key, project.log_public_key)

    def test_the_public_key_goes_into_dist_and_the_private_key_opens_results(self) -> None:
        Frame.objects.create(name="frame-1")
        private_pem, public_pem = generate_key_pair()
        ProjectSettings.objects.update(log_public_key=public_pem)
        generate_dist(ProjectSettings.objects.get())
        self.assertEqual((self.root / "dist" / "log-key.pem").read_text(encoding="utf-8"), public_pem)
        self.assertTemplateUsed(self.client.get(reverse("studio:results")), "studio/unlock.html")
        # a server log, encrypted as log.php writes it, is imported as it is and only readable once unlocked
        event = {"participant_id": "R_1", "visit_id": "v1", "seq": 1, "event": "frame", "frame": "fnr-1", "timestamp": "2026-09-15T08:00:00Z"}
        upload = SimpleUploadedFile("R_1--v1.jsonl", encrypt_event(event, public_pem).encode())
        self.client.post(reverse("studio:import_logs"), {"log_files": upload})
        self.assertFalse((self.root / "project" / "logs" / "R_1--v1.jsonl").exists())
        key_file = SimpleUploadedFile("key.pem", private_pem.encode())
        self.client.post(reverse("studio:unlock_results"), {"key_file": key_file})
        upload = SimpleUploadedFile("R_1--v1.jsonl", encrypt_event(event, public_pem).encode())
        self.client.post(reverse("studio:import_logs"), {"log_files": upload})
        stored = (self.root / "project" / "logs" / "R_1--v1.jsonl").read_text(encoding="utf-8")
        self.assertIn('"enc"', stored)
        page = self.client.get(reverse("studio:results"))
        self.assertTemplateUsed(page, "studio/results.html")
        self.assertContains(page, "R_1")
        self.assertEqual(self.client.get(reverse("studio:log_api", args=["R_1--v1.jsonl"])).json()["events"][0]["frame"], "fnr-1")

    def test_a_key_file_of_another_project_does_not_unlock(self) -> None:
        _private, public_pem = generate_key_pair()
        other_private, _public = generate_key_pair()
        ProjectSettings.objects.update(log_public_key=public_pem)
        self.client.post(reverse("studio:unlock_results"), {"key_file": SimpleUploadedFile("key.pem", other_private.encode())})
        self.assertTemplateUsed(self.client.get(reverse("studio:results")), "studio/unlock.html")

    def test_editing_a_content_row_from_the_frame_editor_writes_the_workbook(self) -> None:
        frame = Frame.objects.create(name="frame-1")
        append_content_row(self.root / "project" / "content.xlsx", "", {"en-US": "Hello"})
        self.client.post(reverse("studio:edit_content_row", args=[1]), {"frame": str(frame.id), "text.en-US": "Hello there", "note": "greeting"})
        row = load_content_table(self.root / "project" / "content.xlsx").by_id()[1]
        self.assertEqual((row.values["en-US"], row.note), ("Hello there", "greeting"))

    def test_results_page_lists_visits_with_readable_events(self) -> None:
        Frame.objects.create(name="Intro")
        frame = Frame.objects.first()
        self.client.post(reverse("studio:preview_log"), json.dumps({**EVENT, "frame": frame.key, "timestamp": "2026-09-10T10:00:00Z"}), content_type="application/json")
        self.assertContains(self.client.get(reverse("studio:results")), "Intro: visited (last frame)")

    def test_project_logo_overrides_the_default(self) -> None:
        (self.root / "project" / "logo-tiltale.png").write_bytes(b"project logo")
        Frame.objects.create(name="frame-1")
        generate_dist(self.project)
        self.assertEqual((self.root / "dist" / "logo-tiltale.png").read_bytes(), b"project logo")

    def test_llm_answer_is_added_to_the_frame(self) -> None:
        frame = Frame.objects.create(name="frame-1")
        target = Frame.objects.create(name="frame-2")
        append_content_row(self.root / "project" / "content.xlsx", "", {"en-US": "Go"})
        answer = ANSWER.replace("fnr-2", target.key)
        self.client.post(reverse("studio:import_elements", kwargs={"frame_id": frame.id}), {"answer": answer})
        self.assertEqual(frame.elements.get().target_frame, target)

    def test_added_image_keeps_its_aspect_ratio(self) -> None:
        Image.new("RGB", (400, 200)).save(self.root / "project" / "materials" / "tree.png")
        frame = Frame.objects.create(name="frame-1")
        self.client.post(reverse("studio:add_image", kwargs={"frame_id": frame.id}), {"image": "tree.png"})
        element = frame.elements.get()
        self.assertEqual((element.image, element.width / element.height), ("tree.png", 2))

    def test_resized_element_keeps_its_new_size(self) -> None:
        element = Element.objects.create(frame=Frame.objects.create(name="frame-1"), component="basic-speech-bubble")
        body = {"x": 100, "y": 200, "width": 300, "height": 150, "language": "en-US"}
        self.client.post(reverse("studio:element_position_api", kwargs={"element_id": element.id}), json.dumps(body), content_type="application/json")
        element.refresh_from_db()
        self.assertEqual((element.x, element.y, element.width, element.height), (100, 200, 300, 150))

    def test_dragged_bubble_tail_is_in_the_story(self) -> None:
        element = Element.objects.create(frame=Frame.objects.create(name="frame-1"), component="basic-speech-bubble")
        body = {"x": 100, "y": 200, "width": 300, "height": 150, "tail_x": -80, "tail_y": 120, "language": "en-US"}
        self.client.post(reverse("studio:element_position_api", kwargs={"element_id": element.id}), json.dumps(body), content_type="application/json")
        generate_dist(self.project)
        story = (self.root / "dist" / "story.js").read_text(encoding="utf-8")
        self.assertIn('"tail_x":-80.0,"tail_y":120.0', story)
        self.assertIn('"basic-speech-bubble":{"svg":', story)
        self.assertTrue((self.root / "dist" / "bubbles.js").is_file())

    def test_moved_background_is_in_the_story(self) -> None:
        frame = Frame.objects.create(name="frame-1", background_type="image", background_image="sky.png")
        body = {"x": 10, "y": 20, "width": 300, "height": 150}
        self.client.post(reverse("studio:background_box_api", kwargs={"frame_id": frame.id}), json.dumps(body), content_type="application/json")
        generate_dist(self.project)
        self.assertIn('"box":{"x":10.0,"y":20.0,"width":300.0,"height":150.0}', (self.root / "dist" / "story.js").read_text(encoding="utf-8"))

    def test_preview_logs_are_encrypted_once_the_project_has_a_key(self) -> None:
        private_pem, public_pem = generate_key_pair()
        ProjectSettings.objects.update(log_public_key=public_pem)
        self.client.post(reverse("studio:preview_log"), json.dumps({**EVENT, "visit_id": "v9"}), content_type="application/json")
        line = (self.root / "project" / "logs" / log_file_name({**EVENT, "visit_id": "v9"})).read_text(encoding="utf-8")
        self.assertTrue(is_encrypted(json.loads(line)))
        private = load_private_key(private_pem.encode(), public_pem)
        self.assertEqual(decrypt_event(json.loads(line), private)["event"], EVENT["event"])
        with override_settings(PLAIN_LOCAL_LOGS=True):
            self.client.post(reverse("studio:preview_log"), json.dumps({**EVENT, "visit_id": "v10"}), content_type="application/json")
        line = (self.root / "project" / "logs" / log_file_name({**EVENT, "visit_id": "v10"})).read_text(encoding="utf-8")
        self.assertFalse(is_encrypted(json.loads(line)))

    def test_preview_log_writes_one_file_per_visit(self) -> None:
        url = reverse("studio:preview_log")
        self.client.post(url, json.dumps(EVENT), content_type="application/json")
        self.client.post(url, json.dumps({**EVENT, "visit_id": "second"}), content_type="application/json")
        names = sorted(path.name for path in (self.root / "project" / "logs").iterdir())
        self.assertEqual(names, ["R_abc--20260910T101530Z-a1b2.jsonl", "R_abc--second.jsonl"])


class MultiLanguageTests(ProjectTestCase):
    languages = ["en-US", "nl-NL"]

    def setUp(self) -> None:
        super().setUp()
        Frame.objects.create(name="frame-1")
        self.picker = Frame.objects.create(name="picker-1", kind="picker")
        Element.objects.create(frame=self.picker, component="basic-decision", text="Nederlands", target_language="nl-NL")

    def test_picker_page_comes_first_then_one_folder_per_language(self) -> None:
        report = generate_dist(self.project)
        self.assertEqual([build.folder for build in report.builds], ["", "en-US", "nl-NL"])

    def test_every_language_folder_has_its_own_page(self) -> None:
        generate_dist(self.project)
        self.assertTrue((self.root / "dist" / "nl-NL" / "index.html").is_file())

    def test_picker_frames_are_not_in_the_language_pages(self) -> None:
        generate_dist(self.project)
        story_js = (self.root / "dist" / "en-US" / "story.js").read_text(encoding="utf-8")
        self.assertNotIn(f'"{self.picker.key}"', story_js)

    def test_a_language_without_a_picker_button_is_an_error(self) -> None:
        issues = validate_project(self.project, load_content_table(self.root / "project" / "content.xlsx"))
        self.assertIn("No language-picker button leads to en-US.", [issue.message for issue in issues])

    def test_picker_frames_alone_do_not_count_as_story_frames(self) -> None:
        Frame.objects.exclude(kind="picker").delete()
        issues = validate_project(self.project, load_content_table(self.root / "project" / "content.xlsx"))
        self.assertIn("Picker frames only make the start page", issues[-1].message)

    def test_tidy_up_puts_picker_frames_left_of_the_story(self) -> None:
        tidy_layout()
        self.assertLess(Frame.objects.get(name="picker-1").flow_x, Frame.objects.get(name="frame-1").flow_x)


class ContentPageTests(ProjectTestCase):
    """The Content page: texts, images and fonts without opening the source files."""

    def test_a_row_can_be_changed_and_added_without_a_reload(self) -> None:
        append_content_row(settings.PROJECT_DIR / "content.xlsx", "greet", {"en-US": "Hello"})
        body = {"id": 1, "note": "greet", "values": {"en-US": "Hi"}}
        self.assertEqual(self.client.post(reverse("studio:content_row_api"), json.dumps(body), content_type="application/json").status_code, 200)
        added = self.client.post(reverse("studio:content_row_api"), json.dumps({"note": "", "values": {"en-US": "Bye"}}), content_type="application/json")
        self.assertEqual(added.json()["id"], 2)
        table = load_content_table(settings.PROJECT_DIR / "content.xlsx")
        self.assertEqual([row.values["en-US"] for row in table.rows], ["Hi", "Bye"])

    def test_an_image_is_uploaded_with_a_safe_name_and_protected_while_used(self) -> None:
        from io import BytesIO
        buffer = BytesIO()
        Image.new("RGB", (8, 8)).save(buffer, "PNG")
        self.client.post(reverse("studio:upload_materials"), {"images": SimpleUploadedFile("My Photo!.png", buffer.getvalue())})
        self.assertEqual(list_materials(), ["My-Photo.png"])
        Frame.objects.create(name="frame-1", background_type="image", background_image="My-Photo.png")
        self.client.post(reverse("studio:delete_material"), {"name": "My-Photo.png"})
        self.assertEqual(list_materials(), ["My-Photo.png"])  # used and not confirmed: kept
        self.client.post(reverse("studio:delete_material"), {"name": "My-Photo.png", "force": "1"})
        self.assertEqual(list_materials(), [])

    def test_choosing_a_font_writes_only_known_stacks_to_the_css(self) -> None:
        body = {"component": "basic-narrator", "preset": "Serif"}
        self.assertEqual(self.client.post(reverse("studio:save_component_font"), json.dumps(body), content_type="application/json").status_code, 200)
        css = (settings.PROJECT_DIR / "style-overrides.css").read_text(encoding="utf-8")
        self.assertIn('.component-basic-narrator { font-family: Georgia, "Times New Roman", serif; }', css)
        bad = {"component": "basic-narrator", "preset": "evil; } body { display: none"}
        self.assertEqual(self.client.post(reverse("studio:save_component_font"), json.dumps(bad), content_type="application/json").status_code, 400)

class OfflineUseTests(ProjectTestCase):
    def test_regenerate_writes_the_offline_page_and_service_worker(self) -> None:
        Frame.objects.create(name="frame-1")
        generate_dist(self.project)
        page = (self.root / "dist" / "offline" / "index.html").read_text(encoding="utf-8")
        self.assertIn('"tiltale:demo:unsent"', page)
        self.assertIn('"folder": ""', page)  # the single page: the story at the dist root
        self.assertIn("networkFirst", (self.root / "dist" / "sw.js").read_text(encoding="utf-8"))

    def test_regenerate_writes_the_reset_page_for_this_project(self) -> None:
        Frame.objects.create(name="frame-1")
        generate_dist(self.project)
        page = (self.root / "dist" / "reset" / "index.html").read_text(encoding="utf-8")
        self.assertIn('"tiltale:demo:"', page)
        self.assertNotIn("__PROJECT__", page)

    def test_a_replaced_image_with_the_same_name_gets_a_new_address(self) -> None:
        photo = self.root / "project" / "materials" / "photo.png"
        Image.new("RGB", (8, 8), "red").save(photo)
        Frame.objects.create(name="frame-1", background_type="image", background_image="photo.png")
        generate_dist(self.project)
        before = sorted(path.name for path in (self.root / "dist" / "assets").iterdir())
        Image.new("RGB", (8, 8), "blue").save(photo)
        generate_dist(self.project)
        after = sorted(path.name for path in (self.root / "dist" / "assets").iterdir())
        self.assertNotEqual(before, after)

class DocumentPageTests(TestCase):
    """The Docs pages: repository documents rendered inside the studio."""

    databases = {"project"}

    def test_the_readme_is_rendered_as_a_page(self) -> None:
        response = self.client.get(reverse("studio:document", kwargs={"slug": "readme"}))
        self.assertContains(response, "<h1>TilTale</h1>")

    def test_unknown_documents_are_not_served(self) -> None:
        self.assertEqual(self.client.get(reverse("studio:document", kwargs={"slug": "settings"})).status_code, 404)


class LanguageColumnTests(SimpleTestCase):
    def test_a_language_column_can_be_added_and_renamed(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "content.xlsx"
            create_content_workbook(path, ["en-US"])
            add_language(path, "nl-NL")
            rename_language(path, "nl-NL", "nl-BE")
            self.assertEqual(load_content_table(path).languages, ("en-US", "nl-BE"))
            with self.assertRaisesRegex(ValueError, "already exists"):
                add_language(path, "NL-be")
            with self.assertRaisesRegex(ValueError, "no language column"):
                rename_language(path, "de-DE", "de-AT")


class SaveLanguagesViewTests(ProjectTestCase):
    languages = ["en-US", "nl-NL"]

    def test_renaming_a_language_updates_everything_that_stores_the_code(self) -> None:
        frame = Frame.objects.create(name="frame-1")
        element = Element.objects.create(frame=frame)
        ElementLanguageOverride.objects.create(element=element, language="nl-NL", x=1)
        picker = Element.objects.create(frame=Frame.objects.create(name="picker-1", kind="picker"), target_language="nl-NL")
        self.client.post(reverse("studio:save_languages"), {"rename.en-US": "en-US", "rename.nl-NL": "nl-BE"})
        self.assertEqual(load_content_table(settings.PROJECT_DIR / "content.xlsx").languages, ("en-US", "nl-BE"))
        self.assertEqual(ElementLanguageOverride.objects.get().language, "nl-BE")
        self.assertEqual(Element.objects.get(pk=picker.pk).target_language, "nl-BE")

    def test_adding_a_language_normalizes_the_code(self) -> None:
        self.client.post(reverse("studio:save_languages"), {"rename.en-US": "en-US", "rename.nl-NL": "nl-NL", "new_language": " de_de "})
        self.assertEqual(load_content_table(settings.PROJECT_DIR / "content.xlsx").languages, ("en-US", "nl-NL", "de-DE"))


class MaterialsPageTests(ProjectTestCase):
    def test_thumbnails_are_small_cached_webp_files(self) -> None:
        Image.new("RGB", (4000, 3000)).save(settings.PROJECT_DIR / "materials" / "big.png")
        thumbnail = material_thumbnail("big.png")
        with Image.open(thumbnail) as image:
            self.assertLessEqual(max(image.size), 360)
        self.assertEqual(material_thumbnail("big.png"), thumbnail)  # served from the cache
        response = self.client.get(reverse("studio:material_thumb", kwargs={"path": "big.png"}))
        self.assertEqual(response["Content-Type"], "image/webp")
        # A FileResponse holds the thumbnail open until it is consumed; a real server always
        # consumes it, but the test client does not, and Windows cannot delete an open file.
        response.close()

    def test_the_materials_page_lists_size_and_usage(self) -> None:
        Image.new("RGB", (640, 480)).save(settings.PROJECT_DIR / "materials" / "scene.png")
        Frame.objects.create(name="frame-1", background_type="image", background_image="scene.png")
        response = self.client.get(reverse("studio:materials"))
        self.assertContains(response, "640×480")
        self.assertContains(response, "background of frame-1")


class LineBreakAndEssentialsTests(SimpleTestCase):
    def test_a_saved_line_break_shows_as_a_line_break_in_excel(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "content.xlsx"
            create_content_workbook(path, ["en-US"])
            append_content_row(path, "", {"en-US": "one"})
            update_content_row(path, 1, "", {"en-US": "First line\nSecond line"})
            cell = load_workbook(path).active.cell(2, 3)
        self.assertEqual((cell.value, cell.alignment.wrap_text), ("First line\nSecond line", True))

    def test_text_just_under_the_minimum_is_not_essential(self) -> None:
        from .services.validate import MIN_TEXT_PX, NEAR_MINIMUM
        self.assertGreater(11.8, MIN_TEXT_PX * NEAR_MINIMUM)  # the reported case is within the 5% margin
        self.assertLess(11.0, MIN_TEXT_PX * NEAR_MINIMUM)


class ExcelOpenTests(ProjectTestCase):
    def test_saving_while_excel_has_the_workbook_open_is_refused_with_a_message(self) -> None:
        append_content_row(settings.PROJECT_DIR / "content.xlsx", "", {"en-US": "Hello"})
        (settings.PROJECT_DIR / "~$content.xlsx").write_bytes(b"owner")  # what Excel leaves while the file is open
        body = {"id": 1, "note": "", "values": {"en-US": "Hi"}}
        response = self.client.post(reverse("studio:content_row_api"), json.dumps(body), content_type="application/json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("open in Excel", response.json()["error"])
        (settings.PROJECT_DIR / "~$content.xlsx").unlink()  # Excel closed: the workbook was left untouched
        self.assertEqual(load_content_table(settings.PROJECT_DIR / "content.xlsx").by_id()[1].values["en-US"], "Hello")
