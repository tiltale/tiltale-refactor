"""Create, read, append to and update rows of the project's ``content.xlsx``.

Layout: column A ``content_id``, column B ``note``, then one column per language.
``content_id`` is a stable integer, deliberately not the Excel row number, so
authors can sort or insert rows freely. A typed row without an ID gets the
next free ID, which is written back to the workbook. Existing IDs never change.
Line breaks inside a cell (Alt+Enter in Excel) are kept and shown as line breaks in the story.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

LOCKED_MESSAGE: str = (
    "content.xlsx is locked. Close it in Excel (or any other program), then reload "
    "this browser page. Restarting TilTale is not needed."
)


@dataclass(frozen=True, slots=True)
class ContentRow:
    content_id: int
    excel_row: int
    values: dict[str, str]
    note: str


@dataclass(frozen=True, slots=True)
class ContentTable:
    languages: tuple[str, ...]
    rows: tuple[ContentRow, ...]

    def by_id(self) -> dict[int, ContentRow]:
        return {row.content_id: row for row in self.rows}


def create_content_workbook(path: Path, languages: list[str]) -> None:
    """Create the empty workbook for a new project."""
    workbook: Workbook = Workbook()
    sheet: Worksheet = workbook.active
    sheet.title = "content"
    sheet.append(["content_id", "note", *languages])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 28
    for index in range(3, 3 + len(languages)):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = 42
    workbook.save(path)


def _text(value: Any) -> str:
    """Trim the ends; keep line breaks inside (browser forms send them as CRLF, Excel as LF)."""
    return "" if value is None else str(value).replace("\r\n", "\n").strip()


def _content_id(value: Any, excel_row: int) -> int:
    """Parse an ID without silently truncating values such as 3.5."""
    error = ValueError(f"Excel row {excel_row}: content_id must be a whole number above zero.")
    if isinstance(value, bool) or (isinstance(value, float) and not value.is_integer()):
        raise error
    try:
        parsed: int = int(str(value).strip()) if not isinstance(value, float) else int(value)
    except ValueError:
        raise error from None
    if parsed <= 0:
        raise error
    return parsed


def _languages(sheet: Worksheet) -> tuple[str, ...]:
    headers: list[str] = [_text(cell.value) for cell in sheet[1]]
    while headers and not headers[-1]:
        headers.pop()
    if headers[:2] != ["content_id", "note"]:
        raise ValueError("content.xlsx must start with the columns 'content_id' and 'note'. Language columns go after them.")
    languages: tuple[str, ...] = tuple(headers[2:])
    if not languages:
        raise ValueError("content.xlsx needs at least one language column after 'note'.")
    if not all(languages):
        raise ValueError("Language columns in content.xlsx must be named; do not leave empty columns between them.")
    if len({language.casefold() for language in languages}) != len(languages):
        raise ValueError("Language column names in content.xlsx must be unique.")
    return languages


def load_content_table(path: Path) -> ContentTable:
    """Read the workbook and assign IDs to newly typed rows."""
    if not path.is_file():
        raise FileNotFoundError(f"Content workbook not found: {path}")
    try:
        workbook = load_workbook(path)
    except PermissionError:
        raise ValueError(LOCKED_MESSAGE) from None
    sheet: Worksheet = workbook.active
    languages: tuple[str, ...] = _languages(sheet)

    # Find the highest existing ID first so a new ID never depends on where
    # an ID-less row happens to sit in the sheet.
    seen: set[int] = set()
    for excel_row in range(2, sheet.max_row + 1):
        raw = sheet.cell(excel_row, 1).value
        if raw in (None, ""):
            continue
        content_id: int = _content_id(raw, excel_row)
        if content_id in seen:
            raise ValueError(f"Duplicate content_id {content_id} in content.xlsx.")
        seen.add(content_id)
    next_id: int = max(seen, default=0) + 1

    rows: list[ContentRow] = []
    changed: bool = False
    for excel_row in range(2, sheet.max_row + 1):
        values: dict[str, str] = {
            language: _text(sheet.cell(excel_row, 3 + index).value) for index, language in enumerate(languages)
        }
        note: str = _text(sheet.cell(excel_row, 2).value)
        raw = sheet.cell(excel_row, 1).value
        if raw in (None, ""):
            if not note and not any(values.values()):
                continue
            sheet.cell(excel_row, 1).value = content_id = next_id
            next_id += 1
            changed = True
        else:
            content_id = _content_id(raw, excel_row)
        rows.append(ContentRow(content_id=content_id, excel_row=excel_row, values=values, note=note))

    if changed:
        try:
            workbook.save(path)
        except PermissionError:
            raise ValueError(LOCKED_MESSAGE) from None
    return ContentTable(languages=languages, rows=tuple(rows))


def _cleaned(table: ContentTable, note: str, values: dict[str, str]) -> tuple[str, dict[str, str]]:
    cleaned: dict[str, str] = {language: _text(values.get(language, "")) for language in table.languages}
    note = _text(note)
    if not note and not any(cleaned.values()):
        raise ValueError("Add a note or text in at least one language.")
    return note, cleaned


def _save(workbook: Workbook, path: Path) -> None:
    try:
        workbook.save(path)
    except PermissionError:
        raise ValueError(LOCKED_MESSAGE) from None


def add_language(path: Path, code: str) -> None:
    """Add an empty language column to the workbook (the header row decides the languages)."""
    table: ContentTable = load_content_table(path)
    if code.casefold() in {language.casefold() for language in table.languages}:
        raise ValueError(f"The language {code} already exists.")
    workbook = load_workbook(path)
    sheet: Worksheet = workbook.active
    column: int = 3 + len(table.languages)
    sheet.cell(1, column).value = code
    sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 42
    _save(workbook, path)


def rename_language(path: Path, old: str, new: str) -> None:
    """Rename a language column header; the texts in the column stay."""
    table: ContentTable = load_content_table(path)
    if old not in table.languages:
        raise ValueError(f"There is no language column {old}.")
    if new.casefold() in {language.casefold() for language in table.languages if language != old}:
        raise ValueError(f"The language {new} already exists.")
    workbook = load_workbook(path)
    workbook.active.cell(1, 3 + table.languages.index(old)).value = new
    _save(workbook, path)


def append_content_row(path: Path, note: str, values: dict[str, str]) -> ContentRow:
    """Append one author-created row and return it with its new stable ID."""
    table: ContentTable = load_content_table(path)
    note, cleaned = _cleaned(table, note, values)
    next_id: int = max((row.content_id for row in table.rows), default=0) + 1
    workbook = load_workbook(path)
    sheet: Worksheet = workbook.active
    sheet.append([next_id, note, *cleaned.values()])
    _save(workbook, path)
    return ContentRow(content_id=next_id, excel_row=sheet.max_row, values=cleaned, note=note)


def update_content_row(path: Path, content_id: int, note: str, values: dict[str, str]) -> ContentRow:
    """Overwrite the note and texts of an existing row; its ID and Excel row stay."""
    table: ContentTable = load_content_table(path)
    row: ContentRow | None = table.by_id().get(content_id)
    if row is None:
        raise ValueError(f"content_id {content_id} is not in content.xlsx.")
    note, cleaned = _cleaned(table, note, values)
    workbook = load_workbook(path)
    sheet: Worksheet = workbook.active
    for column, value in enumerate([note, *cleaned.values()], start=2):
        sheet.cell(row.excel_row, column).value = value
    _save(workbook, path)
    return ContentRow(content_id=content_id, excel_row=row.excel_row, values=cleaned, note=note)
